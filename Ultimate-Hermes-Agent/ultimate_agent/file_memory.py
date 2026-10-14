from __future__ import annotations

import hashlib
import os
import sqlite3
import zipfile
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Iterator

from .config import MEMORY_DATABASE_PATH, ensure_directories

TEXT_EXTENSIONS = {
    ".txt", ".md", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".log", ".py", ".js", ".jsx", ".ts", ".tsx", ".html",
    ".css", ".scss", ".sql", ".sh", ".bat", ".ps1", ".xml", ".html", ".vue", ".svelte",
    ".go", ".rs", ".java", ".kt", ".c", ".cpp", ".h", ".hpp", ".cs", ".swift", ".ipynb",
}
SKIP_DIRS = {".git", ".venv", "venv", "node_modules", "__pycache__", "dist", "build", ".next", "target", "vendor"}
SENSITIVE_FILENAMES = {"credentials.json", "secrets.json", "service-account.json", "firebase-applet-config.json", "id_rsa", "id_ed25519"}
SENSITIVE_SUFFIXES = {".pem", ".p12", ".pfx", ".key", ".keystore"}
MAX_FILE_BYTES = 6 * 1024 * 1024
MAX_FILES_PER_SCAN = 12000
CHUNK_SIZE = 2200
CHUNK_OVERLAP = 250


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        raw = path.read_bytes()
        if b"\x00" in raw[:4096]:
            return ""
        return raw.decode("utf-8", errors="replace")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages[:250])
        except Exception:
            return ""
    if suffix == ".docx":
        try:
            from docx import Document
            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            return "\n".join(parts)
        except Exception:
            return ""
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets[:20]:
                rows.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(max_row=500, values_only=True):
                    vals = [str(v) if v is not None else "" for v in row]
                    if any(vals): rows.append(" | ".join(vals))
            wb.close()
            return "\n".join(rows)
        except Exception:
            return ""
    return ""


def _chunks(text: str) -> list[str]:
    text = text.replace("\x00", "").strip()
    if not text:
        return []
    result = []
    start = 0
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        result.append(text[start:end])
        if end == len(text): break
        start = max(start + 1, end - CHUNK_OVERLAP)
    return result


class FileMemory:
    def __init__(self, path: Path = MEMORY_DATABASE_PATH) -> None:
        ensure_directories()
        self.path = path
        with self._connect() as con:
            con.executescript("""
            CREATE TABLE IF NOT EXISTS indexed_files (
                path TEXT PRIMARY KEY, size INTEGER NOT NULL, mtime REAL NOT NULL,
                digest TEXT NOT NULL, indexed_at TEXT NOT NULL, chunk_count INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS chunks (
                id INTEGER PRIMARY KEY AUTOINCREMENT, path TEXT NOT NULL, chunk_no INTEGER NOT NULL, text TEXT NOT NULL
            );
            CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(path UNINDEXED, text, tokenize='unicode61');
            CREATE INDEX IF NOT EXISTS idx_chunks_path ON chunks(path);
            """)

    @contextmanager
    def _connect(self) -> Iterator[sqlite3.Connection]:
        con = sqlite3.connect(self.path, timeout=30)
        con.row_factory = sqlite3.Row
        try:
            with con:
                yield con
        finally:
            con.close()

    def index_folder(self, folder: str, recursive: bool = True, force: bool = False) -> dict[str, Any]:
        root = Path(folder).expanduser().resolve()
        if not root.exists() or not root.is_dir():
            raise ValueError("Folder does not exist or is not a directory.")
        scanned = indexed = skipped = unchanged = errors = 0
        with self._connect() as con:
            for base, dirs, files in os.walk(root):
                dirs[:] = [d for d in dirs if d.lower() not in SKIP_DIRS and not d.startswith(".")]
                if not recursive and Path(base) != root:
                    dirs[:] = []
                    continue
                for filename in files:
                    if scanned >= MAX_FILES_PER_SCAN:
                        break
                    lower_name = filename.lower()
                    path = (Path(base) / filename).resolve()
                    # Avoid persisting common credential/config files into the searchable text index.
                    if (lower_name in SENSITIVE_FILENAMES or lower_name == ".env" or lower_name.startswith(".env.")
                            or path.suffix.lower() in SENSITIVE_SUFFIXES):
                        skipped += 1
                        continue
                    if path.suffix.lower() not in TEXT_EXTENSIONS | {".pdf", ".docx", ".xlsx"}:
                        skipped += 1
                        continue
                    scanned += 1
                    try:
                        stat = path.stat()
                        if stat.st_size > MAX_FILE_BYTES:
                            skipped += 1
                            continue
                        old = con.execute("SELECT size,mtime FROM indexed_files WHERE path=?", (str(path),)).fetchone()
                        if old and not force and int(old["size"]) == stat.st_size and float(old["mtime"]) == stat.st_mtime:
                            unchanged += 1
                            continue
                        body = _extract_text(path)
                        chunks = _chunks(body)
                        digest = hashlib.sha256(body.encode("utf-8", errors="replace")).hexdigest()
                        con.execute("DELETE FROM chunks_fts WHERE path=?", (str(path),))
                        con.execute("DELETE FROM chunks WHERE path=?", (str(path),))
                        for idx, chunk in enumerate(chunks):
                            con.execute("INSERT INTO chunks(path,chunk_no,text) VALUES(?,?,?)", (str(path), idx, chunk))
                            con.execute("INSERT INTO chunks_fts(path,text) VALUES(?,?)", (str(path), chunk))
                        con.execute("""INSERT INTO indexed_files(path,size,mtime,digest,indexed_at,chunk_count)
                                     VALUES(?,?,?,?,datetime('now'),?) ON CONFLICT(path) DO UPDATE SET
                                     size=excluded.size,mtime=excluded.mtime,digest=excluded.digest,
                                     indexed_at=excluded.indexed_at,chunk_count=excluded.chunk_count""",
                                     (str(path), stat.st_size, stat.st_mtime, digest, len(chunks)))
                        indexed += 1
                    except Exception:
                        errors += 1
                if scanned >= MAX_FILES_PER_SCAN:
                    break
        return {"folder": str(root), "scanned": scanned, "indexed_or_updated": indexed,
                "unchanged": unchanged, "skipped": skipped, "errors": errors,
                "max_files_per_scan": MAX_FILES_PER_SCAN}

    def search(self, query: str, limit: int = 8) -> list[dict[str, Any]]:
        query = query.strip()
        if not query:
            return []
        limit = max(1, min(int(limit), 20))
        with self._connect() as con:
            try:
                terms = " AND ".join('"' + part.replace('"', '""') + '"' for part in query.split()[:12])
                rows = con.execute("""SELECT path,text,bm25(chunks_fts) AS rank FROM chunks_fts
                                  WHERE chunks_fts MATCH ? ORDER BY rank LIMIT ?""", (terms, limit)).fetchall()
            except sqlite3.OperationalError:
                like = f"%{query[:100]}%"
                rows = con.execute("SELECT path,text,0 AS rank FROM chunks_fts WHERE text LIKE ? LIMIT ?",
                                   (like, limit)).fetchall()
        return [{"path": row["path"], "excerpt": row["text"][:1400], "rank": float(row["rank"])} for row in rows]

    def stats(self) -> dict[str, int]:
        with self._connect() as con:
            files = con.execute("SELECT COUNT(*) AS n FROM indexed_files").fetchone()["n"]
            chunks = con.execute("SELECT COUNT(*) AS n FROM chunks").fetchone()["n"]
        return {"indexed_files": int(files), "chunks": int(chunks)}

    def analyze_archive(self, archive_path: str) -> dict[str, Any]:
        path = Path(archive_path).expanduser().resolve()
        if not path.is_file() or path.suffix.lower() != ".zip":
            raise ValueError("Please provide the path to a .zip file.")
        max_entries, max_uncompressed = 30000, 700 * 1024 * 1024
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
            if len(infos) > max_entries:
                infos = infos[:max_entries]
            total = 0
            roots: dict[str, int] = {}
            notable: list[str] = []
            readmes: list[dict[str, str]] = []
            source_exts: dict[str, int] = {}
            for info in infos:
                if info.is_dir(): continue
                total += max(0, info.file_size)
                if total > max_uncompressed: break
                name = info.filename.replace("\\", "/")
                parts = [part for part in name.split("/") if part]
                if parts:
                    roots[parts[0]] = roots.get(parts[0], 0) + 1
                suffix = Path(name).suffix.lower() or "[no extension]"
                source_exts[suffix] = source_exts.get(suffix, 0) + 1
                low = name.lower()
                if any(k in low for k in ("readme", "features.md", "skills/", "tools/", "tool/", "hotkey", "shortcut", "manifest", "pyproject.toml", "package.json", "requirements.txt")):
                    if len(notable) < 200: notable.append(name)
                base = Path(name).name.lower()
                if base.startswith("readme") and len(readmes) < 12 and info.file_size < 80000:
                    try:
                        content = zf.read(info).decode("utf-8", errors="replace")[:5000]
                        readmes.append({"path": name, "preview": content})
                    except Exception:
                        pass
            return {"archive": str(path), "entries_considered": min(len(zf.infolist()), max_entries),
                    "uncompressed_bytes_considered": min(total, max_uncompressed),
                    "top_level_roots": sorted(roots.items(), key=lambda x: x[1], reverse=True)[:60],
                    "file_extensions": sorted(source_exts.items(), key=lambda x: x[1], reverse=True)[:30],
                    "notable_files": notable, "readme_previews": readmes,
                    "warning": "Archive paths were inspected without extracting files; content is bounded to reduce zip-bomb risk."}

file_memory = FileMemory()

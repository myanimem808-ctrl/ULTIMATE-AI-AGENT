from __future__ import annotations

import csv
import json
import os
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .config import ROOT, WORKSPACE_DIR

PROTECTED_PATHS = {
    Path(os.environ.get("SystemRoot", "C:\\Windows")).resolve(),
    Path(os.environ.get("ProgramFiles", "C:\\Program Files")).resolve(),
    Path(os.environ.get("ProgramFiles(x86)", "C:\\Program Files (x86)")).resolve(),
    Path("C:\\").resolve(),
}


def _resolve_safe_path(path_str: str) -> Path:
    raw = (path_str or "").strip()
    if not raw or raw.lower() in {".", "here", "current", "workspace"}:
        return WORKSPACE_DIR.resolve()
    if raw.lower() in {"desktop", "my desktop"}:
        return (Path.home() / "Desktop").resolve()
    if raw.lower() in {"documents", "my documents"}:
        return (Path.home() / "Documents").resolve()
    if raw.lower() in {"downloads", "my downloads"}:
        return (Path.home() / "Downloads").resolve()

    p = Path(raw)
    if not p.is_absolute():
        p = WORKSPACE_DIR / p
    return p.resolve()


def _check_not_protected(path: Path) -> None:
    resolved = path.resolve()
    for protected in PROTECTED_PATHS:
        if resolved == protected:
            raise PermissionError(f"Action blocked: '{resolved}' is a protected system directory.")


def create_folder(path: str) -> dict[str, Any]:
    target = _resolve_safe_path(path)
    _check_not_protected(target)
    target.mkdir(parents=True, exist_ok=True)
    return {"ok": True, "action": "create_folder", "path": str(target), "exists": target.is_dir()}


def create_file(path: str, content: str = "") -> dict[str, Any]:
    target = _resolve_safe_path(path)
    _check_not_protected(target)
    if target.exists():
        raise FileExistsError(f"File already exists at: {target}. Use edit_file or write_file instead.")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(content, encoding="utf-8", errors="replace")
    return {"ok": True, "action": "create_file", "path": str(target), "size_bytes": target.stat().st_size}


def read_file(path: str, start_line: int = 1, end_line: int | None = None) -> dict[str, Any]:
    target = _resolve_safe_path(path)
    if not target.is_file():
        raise FileNotFoundError(f"File not found: {target}")
    
    text = target.read_text(encoding="utf-8", errors="replace")
    lines = text.splitlines()
    total_lines = len(lines)
    
    start_idx = max(0, start_line - 1)
    end_idx = total_lines if end_line is None else min(total_lines, max(start_idx, end_line))
    selected_lines = lines[start_idx:end_idx]
    
    formatted_output = "\n".join(f"{start_idx + i + 1:>4} | {line}" for i, line in enumerate(selected_lines))
    return {
        "ok": True,
        "path": str(target),
        "total_lines": total_lines,
        "displayed_range": [start_idx + 1, end_idx],
        "content": formatted_output[:32000],
    }


def write_file(path: str, content: str, overwrite: bool = True) -> dict[str, Any]:
    target = _resolve_safe_path(path)
    _check_not_protected(target)
    if target.exists() and not overwrite:
        raise FileExistsError(f"File already exists: {target}. Set overwrite=True to replace.")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(content, encoding="utf-8", errors="replace")
    return {"ok": True, "action": "write_file", "path": str(target), "size_bytes": target.stat().st_size}


def edit_file(path: str, old_text: str, new_text: str) -> dict[str, Any]:
    target = _resolve_safe_path(path)
    _check_not_protected(target)
    if not target.is_file():
        raise FileNotFoundError(f"File not found: {target}")
    content = target.read_text(encoding="utf-8", errors="replace")
    count = content.count(old_text)
    if count == 0:
        raise ValueError(f"Target string not found in {target.name}.")
    if count > 1:
        raise ValueError(f"Target string appears {count} times in {target.name}. Please provide a unique snippet.")
    new_content = content.replace(old_text, new_text)
    target.write_text(new_content, encoding="utf-8", errors="replace")
    return {"ok": True, "action": "edit_file", "path": str(target), "message": "Text successfully replaced."}


def move_item(source: str, destination: str) -> dict[str, Any]:
    src = _resolve_safe_path(source)
    dst = _resolve_safe_path(destination)
    _check_not_protected(src)
    _check_not_protected(dst)
    if not src.exists():
        raise FileNotFoundError(f"Source not found: {src}")
    dst.parent.mkdir(parents=True, exist_ok=True)
    result = shutil.move(str(src), str(dst))
    return {"ok": True, "action": "move_item", "source": str(src), "destination": str(result)}


def copy_item(source: str, destination: str) -> dict[str, Any]:
    src = _resolve_safe_path(source)
    dst = _resolve_safe_path(destination)
    if not src.exists():
        raise FileNotFoundError(f"Source not found: {src}")
    dst.parent.mkdir(parents=True, exist_ok=True)
    if src.is_dir():
        final_dst = shutil.copytree(str(src), str(dst), dirs_exist_ok=True)
    else:
        final_dst = shutil.copy2(str(src), str(dst))
    return {"ok": True, "action": "copy_item", "source": str(src), "destination": str(final_dst)}


def rename_item(path: str, new_name: str) -> dict[str, Any]:
    target = _resolve_safe_path(path)
    _check_not_protected(target)
    if not target.exists():
        raise FileNotFoundError(f"Path not found: {target}")
    clean_name = Path(new_name).name
    destination = target.parent / clean_name
    target.rename(destination)
    return {"ok": True, "action": "rename_item", "old_path": str(target), "new_path": str(destination)}


def delete_item(path: str) -> dict[str, Any]:
    target = _resolve_safe_path(path)
    _check_not_protected(target)
    if not target.exists():
        raise FileNotFoundError(f"Path not found: {target}")
    if target.is_dir():
        shutil.rmtree(str(target))
    else:
        target.unlink()
    return {"ok": True, "action": "delete_item", "deleted_path": str(target)}


def compress_item(path: str, format: str = "zip", destination: str | None = None) -> dict[str, Any]:
    target = _resolve_safe_path(path)
    if not target.exists():
        raise FileNotFoundError(f"Path not found: {target}")
    dest_dir = _resolve_safe_path(destination) if destination else target.parent
    dest_dir.mkdir(parents=True, exist_ok=True)
    
    fmt = "zip" if "zip" in format.lower() else "tar"
    archive_base = str(dest_dir / target.name)
    if target.is_dir():
        out = shutil.make_archive(archive_base, fmt, str(target))
    else:
        out = f"{archive_base}.zip"
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.write(str(target), arcname=target.name)
    return {"ok": True, "action": "compress_item", "archive_path": out, "size_bytes": os.path.getsize(out)}


def list_directory(path: str = "", sort_by: str = "name", reverse: bool = False) -> dict[str, Any]:
    target = _resolve_safe_path(path)
    if not target.is_dir():
        raise NotADirectoryError(f"Directory not found: {target}")
    
    items = []
    total_size = 0
    for entry in target.iterdir():
        try:
            stat = entry.stat()
            size = stat.st_size if entry.is_file() else 0
            total_size += size
            mtime = datetime.fromtimestamp(stat.st_mtime, timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
            items.append({
                "name": entry.name,
                "path": str(entry),
                "is_dir": entry.is_dir(),
                "size_bytes": size,
                "modified": mtime,
                "extension": entry.suffix.lower() if entry.is_file() else "[DIR]",
            })
        except OSError:
            continue
            
    sort_key = sort_by.lower()
    if "size" in sort_key:
        items.sort(key=lambda x: x["size_bytes"], reverse=reverse or True)
    elif "date" in sort_key or "time" in sort_key:
        items.sort(key=lambda x: x["modified"], reverse=reverse or True)
    elif "type" in sort_key or "ext" in sort_key:
        items.sort(key=lambda x: (not x["is_dir"], x["extension"], x["name"].lower()), reverse=reverse)
    else:
        items.sort(key=lambda x: (not x["is_dir"], x["name"].lower()), reverse=reverse)

    return {
        "ok": True,
        "directory": str(target),
        "total_items": len(items),
        "total_size_bytes": total_size,
        "items": items[:150],
    }


def analyze_document(path: str) -> dict[str, Any]:
    """Deep analysis of documents (.docx, .pdf, .pptx, .xlsx, .csv, .txt, .json, .py)."""
    target = _resolve_safe_path(path)
    if not target.is_file():
        raise FileNotFoundError(f"File not found: {target}")

    ext = target.suffix.lower()
    meta: dict[str, Any] = {"path": str(target), "file_name": target.name, "size_bytes": target.stat().st_size, "extension": ext}
    extracted_text = ""

    if ext == ".docx":
        import docx
        doc = docx.Document(str(target))
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        tables_data = []
        for table in doc.tables[:5]:
            rows = []
            for row in table.rows[:10]:
                rows.append([c.text.strip() for c in row.cells])
            if rows:
                tables_data.append(rows)
        meta["paragraphs_count"] = len(doc.paragraphs)
        meta["tables_count"] = len(doc.tables)
        extracted_text = "\n".join(paras[:100])
        if tables_data:
            extracted_text += "\n\n[Tables Summary]:\n" + json.dumps(tables_data, indent=2)

    elif ext == ".pdf":
        import pypdf
        reader = pypdf.PdfReader(str(target))
        meta["pages_count"] = len(reader.pages)
        page_texts = []
        for i, page in enumerate(reader.pages[:20]):
            txt = page.extract_text() or ""
            if txt.strip():
                page_texts.append(f"--- Page {i+1} ---\n{txt}")
        extracted_text = "\n\n".join(page_texts)

    elif ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(target))
        meta["slides_count"] = len(prs.slides)
        slide_texts = []
        for i, slide in enumerate(prs.slides[:30]):
            shapes_text = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    shapes_text.append(shape.text.strip())
            slide_texts.append(f"--- Slide {i+1} ---\n" + "\n".join(shapes_text))
        extracted_text = "\n\n".join(slide_texts)

    elif ext in {".xlsx", ".xls"}:
        import openpyxl
        wb = openpyxl.load_workbook(str(target), data_only=True)
        meta["sheet_names"] = wb.sheetnames
        sheets_data = []
        for name in wb.sheetnames[:4]:
            ws = wb[name]
            sheet_rows = []
            for row in ws.iter_rows(max_row=20, max_col=10, values_only=True):
                if any(row):
                    sheet_rows.append([str(c) if c is not None else "" for c in row])
            sheets_data.append({"sheet": name, "rows": sheet_rows})
        extracted_text = json.dumps(sheets_data, indent=2)

    elif ext == ".csv":
        with open(target, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            rows = [r for _, r in zip(range(40), reader)]
            meta["preview_rows"] = len(rows)
            extracted_text = "\n".join(", ".join(r) for r in rows)

    else:
        extracted_text = target.read_text(encoding="utf-8", errors="replace")[:25000]

    return {
        "ok": True,
        "metadata": meta,
        "content_excerpt": extracted_text[:16000],
    }
